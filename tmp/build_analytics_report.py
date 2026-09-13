from pathlib import Path
import json
import hashlib
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
S = json.loads((ROOT / 'backend/simulation/output/hybrid-evaluation/scores.json').read_text())
P = json.loads((ROOT / 'backend/simulation/output/paper-based/results.json').read_text())
E = json.loads((ROOT / 'backend/simulation/output/validation-report.json').read_text())
OUT = ROOT / 'docs/BloodConnect_Analytics_Testing_Report.docx'
doc = Document()
# Preset: standard_business_brief. Header: memo_masthead, no border.
# Named overrides: TableBody 10pt/1.05/3pt after; Small 9pt; Title 23pt.
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(.492)

def style(name, size, color='222222', before=0, after=6, line=1.10, bold=False):
    st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name, st.font.size, st.font.bold = 'Calibri', Pt(size), bold
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
    st.paragraph_format.line_spacing = line
    st.paragraph_format.widow_control = True
    return st
style('Normal', 11)
style('Title', 23, '0B2545', after=4, bold=True)
style('Subtitle', 13, '555555', after=12)
for name, size, color, before, after in [('Heading 1',16,'2E74B5',16,8),('Heading 2',13,'2E74B5',12,6),('Heading 3',12,'1F4D78',8,4)]:
    style(name,size,color,before,after,bold=True).paragraph_format.keep_with_next = True
style('TableBody',10,after=3,line=1.05)
style('TableHead',10,after=3,line=1.05,bold=True)
style('Small',9,'555555',after=4)
style('TableSource',9,'555555',before=4,after=4)
style('Caution',11,'7A5A00',before=6,after=8,bold=True)
style('Header',9,'555555',after=0)
style('Footer',9,'555555',after=0)
section.header.paragraphs[0].text = 'BloodConnect | Analytics testing record'
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Exploratory simulation | Page ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title = 'BloodConnect Predictive and Prescriptive Analytics Testing Report'
doc.core_properties.subject = 'Controlled tests and published-mean-informed synthetic evaluation'
doc.core_properties.author = 'BloodConnect project'

def p(text, st=None): return doc.add_paragraph(text, st)
def h(text, level=1): return doc.add_heading(text, level)
def page(title): doc.add_page_break(); h(title)
def f(value): return 'N/A' if value is None else f'{value:.2f}'
def pct(value): return f(value) + '%'
def table(headers, rows, widths):
    assert sum(widths) == 9360
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    pr = t._tbl.tblPr
    w = pr.find(qn('w:tblW')); w.set(qn('w:w'),'9360'); w.set(qn('w:type'),'dxa')
    ind = OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); pr.append(ind)
    margins = OxmlElement('w:tblCellMar')
    for name, val in [('top',80),('bottom',80),('start',120),('end',120)]:
        elem = OxmlElement('w:'+name); elem.set(qn('w:w'),str(val)); elem.set(qn('w:type'),'dxa'); margins.append(elem)
    pr.append(margins)
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D7DBE2'); borders.append(b)
    pr.append(borders)
    grid = t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'),str(width)); grid.append(col)
    for values in rows: t.add_row()
    for i, values in enumerate([headers]+rows):
        row = t.rows[i]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if i == 0: trPr.append(OxmlElement('w:tblHeader'))
        for j, val in enumerate(values):
            cell = row.cells[j]; cell.width = Inches(widths[j]/1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(widths[j]))
            para = cell.paragraphs[0]; para.style = doc.styles['TableHead' if i == 0 else 'TableBody']; para.text = str(val)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                fill = OxmlElement('w:shd'); fill.set(qn('w:fill'),'F2F4F7'); cell._tc.get_or_add_tcPr().append(fill)
    return t

d=S['demand']; s=S['shortageNoReplenishment']; r=S['prescriptive']; b=S['baselineLastWeek']; sr=S['shortageWithDelayedReceipts']
p('ANALYTICS TESTING REPORT','Title')
p('BloodConnect | Predictive and prescriptive analytics','Subtitle')
p('Prepared for: Capstone research documentation\nReport date: 8 September 2026\nEvidence: Saved test artifacts and evaluator source code','Small')
h('Executive summary')
p('The current analytics engine was tested in three stages: an initial small controlled fixture, paper-based calculation checks, and a larger seeded synthetic evaluation. The latest evaluation generated 29,700 daily records, 2,805 seven-day forecasts, and 540 transfer-plan scenarios containing 630 pending requests.')
table(['Module','Latest measure','Score'],[
 ['Demand and usage','Forecasts within +/-20% of simulated usage',pct(d['within20PercentScore'])],
 ['Shortage forecast','Binary accuracy; no replenishment',pct(s['accuracyPercent'])],
 ['Transfer recommendations','Plans satisfying all tested constraints',pct(r['feasiblePlanPercent'])],
], [2500,5000,1860])
p('Sources: hybrid-evaluation/scores.json; 1,552/2,805 tolerance hits, 2,297/2,805 correct shortage labels, 270/540 feasible transfer plans.','TableSource')
p('These are exploratory simulation results, not estimates of real-world hospital accuracy. There is no valid single combined accuracy score.','Caution')
h('What the results mean',2)
p(f"Demand forecasts had MAE {f(d['MAE'])} units per seven-day forecast and WAPE {pct(d['WAPE_percent'])}. Shortage F1 was {pct(s['F1Percent'])}; shortage accuracy fell to {pct(sr['accuracyPercent'])} when simulated receipts and delays were replayed. Three known recommendation failure types remained unfixed.")
p('The synthetic generation was anchored to published average usage levels only. Variability, temporal behavior, demand shocks, stock coverage, expiry and delivery assumptions were not calibrated to original hospital records. The papers do not provide those daily datasets.')
h('Scope',2)
p('Direct tests exercised calculateUsageTrends, calculateShortageForecast and calculateTransferRecommendations in frontend/src/admin/analyticsEngine.js. They did not validate browser flows, deployed APIs, all database behavior, clinical decisions, or donor-selection functions. The latest evaluation did not change the engine or access the production database.')

page('Testing history and score reconciliation')
h('Stage 1: Initial small controlled fixture',2)
p(f"The retained validation-report.json records {E['passed']}/{E['checks']} passing checks. Its four controlled future-usage comparisons gave MAE 0.125 units and MAPE 8.33%. The previously quoted 91.67% was simply 100% minus MAPE; it was not a general system accuracy or real-world result. This document reviews that saved run, rather than claiming it was rerun during document preparation.")
table(['Group','Predicted 7-day units','Controlled outcome'],[['O+',14,14],['A+',7,7],['AB+','2.3333',2],['O-','1.1667',1]],[2400,3480,3480])
p('Source: initial validation-report.json. All four entries concern whole-blood controlled test cases.','TableSource')
h('Stage 2: Published-aggregate calculation tests',2)
p('Seven India phase means and four Korean ABO means were converted into artificial constant-rate histories. These checked arithmetic and rule execution, not predictions against independently observed future hospital usage.')
table(['Category','Passed','Total'],[['Demand/usage arithmetic',11,11],['Shortage arithmetic',11,11],['Constructed recommendation scenarios',3,6],['Red-cell compatibility',0,2]],[6200,1580,1580])
p('Source: paper-based/results.json. The reported 89.29% was 25/28 pass rate excluding compatibility; 83.33% was 25/30 including it. Neither was forecasting accuracy.','TableSource')
h('Stage 3: Independent synthetic outcome evaluation',2)
p('The later seeded benchmark replaced constant daily inputs with variable integer observations and held-out seven-day outcomes. It scored demand error, shortage detection and transfer-plan feasibility separately. Its 50% plan-feasibility result is based on 540 scenarios, not merely the earlier six cases. It does not supersede product-compatibility failures.')
p('A 100% formula-check pass rate can coexist with imperfect forecasting: arithmetic correctness and uncertainty about future demand are different evaluation questions.')

page('Published sources and extracted data')
h('North India study [1]',2)
p('Blood Inventory Management During COVID-19 Pandemic Using a Simple Mathematical Tool: A Two-Year Study from a Tertiary Care Hospital in North India. The study covers 1 October 2019 to 31 December 2021. Table 1 (PDF page 3; printed page 657) gives seven unequal phase summaries, combining packed red cells and whole blood.')
table(['Phase','Collected/week','Issued/week','Average stock','Safety stock'],[
 [x['phase'],x['weeklyCollection'],x['weeklyIssued'],x['averageStock'],x['safetyStock']] for x in P['publishedAggregates']['india']
],[1120,2060,2060,2060,2060])
p('Published mean values, in units. Standard deviations were not used to calibrate the synthetic generator. Stage 2 used mean stock as hypothetical snapshots; Stage 3 used only issued/week divided by seven as rate anchors.','TableSource')
h('Korean hospital study [2]',2)
p('Optimizing the Hospital Blood Bank Stock in Korea: A Comparative Analysis of the Uniform 5-Day Stock Index and a Novel Blood Stock Index. The study used 2019-2020 RBC records from 194 hospitals. The extracted values are from Table 1, Total (N=194), PDF page 5; printed page 266.')
table(['ABO','Mean daily usage','Mean stock index (days)'],[
 [x['abo'],x['averageDailyUsage'],x['averageStockIndexDays']] for x in P['publishedAggregates']['korea']
],[1600,3480,4280])
p('Daily usage includes transfused and wasted units. RhD is unspecified. Stage 3 used daily means as numerical scale references; it did not use the mean stock indices to infer actual stock.','TableSource')
p('The papers supply aggregate evidence, not the original daily files. Generated whole-blood records are supported-component test surrogates, not faithful conversions of the papers\' RBC records. No actual ABO/Rh distribution is reconstructed.','Caution')

page('Synthetic generation: relationships and assumptions')
p('The generator is separate from the analytics engine and never uses model predictions to generate future outcomes. It uses eleven nominal rate anchors (seven India phases and four Korean ABO groups), five pattern families and three fixed seeds. This produces 165 separate 180-day series. Fixed seeds reproduce the same records; they do not make the assumptions clinically representative.')
table(['Pattern','Exact generation rule'],[
 ['Stationary','Poisson daily counts at the nominal rate.'],
 ['Weekday cycle','Rate multiplied by Sun-Sat factors: 0.55, 1.15, 1.20, 1.20, 1.15, 1.10, 0.65.'],
 ['Demand spikes','10% daily chance of tripling the nominal Poisson rate.'],
 ['Level change','Rate factors: 1 before day 90; 1.70 on days 90-134; 0.65 from day 135 (zero-based).'],
 ['Intermittent','55% daily chance of zero usage; otherwise Poisson rate = nominal rate / 0.45.'],
],[2500,6860])
p('All distribution choices above are test assumptions, not estimates from hospital microdata. Poisson variability was not fitted to the papers\' reported standard deviations.','TableSource')
h('Inventory, expiry and receipts',2)
p('At each forecast origin, independently sampled stock equals the nominal daily rate times 2, 5, 8 or 14 days of coverage, times a uniform factor from 0.8 to 1.2, rounded to an integer. Up to 45% is assigned to a batch expiring in 1-6 days; the remainder expires in 30 days. These snapshots do not form one continuous hospital ledger.')
p('A receipt is generated with 60% probability. Its planned day is 1-4 days after origin; a three-day delay has 35% probability. Receipt size is a rounded 2-5 nominal days of usage. Expiry and arrivals are generated assumptions and are recorded explicitly.')
h('Requests and locations',2)
p('Six prescriptive families each contain 90 scenarios: routine transfer, already covered, competing requests, stale expired status, fragmented supply and no supply. Request size is rounded nominal weekly usage times a uniform factor from 0.5 to 1.5. Four synthetic hospitals and optional central stock are used. Priorities and source quantities are constructed per family.')
p('Seeds: 104729, 130363, 155921. No seeds or parameters were selected to improve the scores. India uses ALL_UNSPECIFIED; Korean cases use unsigned ABO. All component values in this benchmark are synthetic whole_blood.')

page('Evaluation procedure and metric definitions')
h('Chronological evaluation',2)
p('Each series starts with 60 days of history. The first forecast origin is day 60; origins advance by seven days while a full next-week outcome remains. This gives 17 origins per series and 2,805 comparisons. Forecast time is midnight, and historical requests are dated noon on earlier days. Only prior records are passed to the engine. Future test horizons do not overlap within a series; historical windows do overlap.')
p('Later forecasts can use previously evaluated days once they have become history. Synthetic historical fulfilled usage equals generated demand by assumption: no historical supply censoring is modeled. This is rolling-origin evaluation of the existing rule-based engine, not training a new machine-learning model.')
h('Engine behavior under test',2)
p('Demand estimate = fulfilled usage in the preceding 30 days / 30 x 7. Historical trend compares current and previous 30-day usage, with increasing/decreasing beyond +/-10%. Shortage days = usable stock / recent daily usage. The engine removes all units expiring within seven days when estimating usable stock; it does not model planned replenishment. Critical is below seven days, low below fourteen.')
h('Demand measures [3]',2)
p('For predicted value f and simulated actual y: MAE = mean(|f-y|); RMSE = sqrt(mean((f-y)^2)); WAPE = 100 x sum(|f-y|)/sum(y); MAPE = 100 x mean(|f-y|/y) over nonzero y only. Bias = mean(f-y). Lower error is better; positive bias means overprediction.')
p('Tolerance score = percentage with |f-y| <= 0.20 x y. A zero outcome qualifies only when the forecast is also zero. The +/-20% threshold was declared for this benchmark, not established as a clinical acceptance standard. Seven zero-outcome weeks are excluded only from MAPE.')
h('Shortage and recommendation measures [4]',2)
p('Shortage accuracy = (TP+TN)/N; precision = TP/(TP+FP); recall = TP/(TP+FN); F1 = 2TP/(2TP+FP+FN). Balanced accuracy averages recall and specificity. Transfer-plan feasibility = plans satisfying every tested constraint / all scenarios. These measures cannot be averaged into one meaningful system accuracy.')
p('A last-week-repeat demand baseline, a majority-class shortage baseline and a separately implemented feasible greedy transfer comparator provide context. No clinician-approved acceptance threshold or expert recommendation oracle was supplied.')

page('Demand and usage results')
table(['Measure','BloodConnect','Last-week baseline'],[
 ['MAE (units / seven-day forecast)',f(d['MAE']),f(b['MAE'])],
 ['RMSE (units / seven-day forecast)',f(d['RMSE']),f(b['RMSE'])],
 ['WAPE',pct(d['WAPE_percent']),pct(b['WAPE_percent'])],
 ['MAPE (nonzero outcomes)',pct(d['MAPE_nonzero_percent']),pct(b['MAPE_nonzero_percent'])],
 ['Within +/-20%',pct(d['within20PercentScore']),pct(b['within20PercentScore'])],
 ['Mean bias (units)',f(d['bias']),f(b['bias'])],
],[5200,2080,2080])
p('Source: scores.json. 2,805 forecasts; MAPE includes 2,798 nonzero outcomes and excludes seven zero outcomes. Baseline forecast equals the preceding seven days\' usage.','TableSource')
h('Performance by generated pattern',2)
table(['Pattern','MAE','WAPE','Within 20%'],[
 [key.replace('_',' ').title(),f(v['demand']['MAE']),pct(v['demand']['WAPE_percent']),pct(v['demand']['within20PercentScore'])] for key,v in S['byPattern'].items()
],[3600,1600,2080,2080])
p('Each family contains 561 forecast outcomes. Differences reflect the chosen synthetic processes, not measured hospital subgroups.','TableSource')
h('Interpretation',2)
p('The engine outperformed repeating the last week on pooled MAE and WAPE in this benchmark, but only 1,552 of 2,805 forecasts were within the selected 20% tolerance. This does not mean that the remaining forecasts were completely wrong; their errors simply exceeded that tolerance.')
p('Historical trend labels agreed with the independently calculated current-versus-previous usage rule in 2,805/2,805 cases. That 100% agreement is descriptive calculation correctness, not prediction of future direction.')
p('The earlier 100% arithmetic result checked that a constant-rate input reproduced its intended average. The current demand results compare forecasts with separate, variable simulated outcomes; their scores answer a different question.')

page('Blood shortage forecast results')
p('A predicted alert is status critical, critical_out, near_expiry_only or at_risk. Independent ground truth is any unmet generated demand during the next seven days. The replay removes expired batches at the start of each day and consumes earliest-expiring valid units first; it does not use the forecast\'s days-of-stock formula to generate labels.')
table(['Confusion-matrix count','No receipts','With delayed receipts'],[
 ['True positive (alert, shortage)',s['TP'],sr['TP']],['True negative (no alert, no shortage)',s['TN'],sr['TN']],
 ['False positive (alert, no shortage)',s['FP'],sr['FP']],['False negative (no alert, shortage)',s['FN'],sr['FN']],
],[5200,2080,2080])
p('Source: scores.json. Both columns use the same 2,805 engine predictions against different inventory-replay outcomes.','TableSource')
table(['Measure','No receipts','With delayed receipts'],[
 [label,pct(s[key]),pct(sr[key])] for label,key in [
 ('Accuracy','accuracyPercent'),('Precision','precisionPercent'),('Recall','recallPercent'),('F1','F1Percent'),
 ('Specificity','specificityPercent'),('Balanced accuracy','balancedAccuracyPercent'),
 ('Shortage prevalence','prevalencePercent'),('Majority-class baseline accuracy','majorityClassBaselineAccuracyPercent')]
],[5200,2080,2080])
p('The paired tables separate raw classification counts from their derived scores. Values are rounded only for presentation.','TableSource')
h('Interpretation',2)
p('Without receipts, the engine detected 1,522 of 1,574 simulated shortages, but raised 456 alerts where no shortage occurred. High recall therefore coexisted with false alarms. The no-replenishment result is closer to the engine\'s supported assumptions.')
p('When actual simulated receipts were included, some alerts became unnecessary because supply arrived. Accuracy decreased to 71.69%. The engine does not accept planned receipts, so this column is an operational stress test, not evidence of a functioning supply-aware forecasting feature.')
p('These are binary shortage-detection scores. Exact depletion-date error and the accuracy of every multi-level risk label were not evaluated.')

page('Prescriptive recommendation results')
p('A raw transfer plan passes only when it includes the expected priority-ordered requests; uses integer nonnegative quantities; avoids self or incompatible transfers; does not exceed unmet destination need; excludes date-expired supply; and respects cumulative available stock plus a 20-unit reserve at hospital sources. Central stock has no reserve in this test.')
table(['Scenario family','Feasible plans','Feasibility'],[
 [key.replace('_',' ').title(),f"{v['feasiblePlans']}/{v['scenarios']}",pct(v['feasiblePlanPercent'])] for key,v in S['byPrescriptionFamily'].items()
],[5200,2080,2080])
p('Total: 270/540 feasible plans (50%). Six equally weighted stress families were designed; this is not the frequency of safe recommendations in routine practice.','TableSource')
h('Observed failure types',2)
p('Already covered: a request can display an already-covered message while still showing positive suggested transfer units. Competing requests: independent recommendations can reuse the same sendable source units. Stale expiry status: a past expiration date is not rejected when status still says available. These failures remain in the tested engine.')
p('The earlier concrete cases showed a 39-unit suggestion for an already-covered request; 78 combined suggested units against only 56 sendable; and a 39-unit suggestion from date-expired stock. Both tested red-cell labels also normalized to whole_blood in separate compatibility probes.')
h('Fulfillment is a different score',2)
p(f"The external evaluator capped or rejected unsafe suggestions before calculating hypothetical fulfillment: {r['fulfilledWithSafetyCaps']:,}/{r['totalDemand']:,} units, or {pct(r['fulfillmentPercentWithSafetyCaps'])}. No transfer filled {pct(r['noTransferBaselineFulfillmentPercent'])}; the feasible split-source greedy comparator filled {pct(r['greedyReferenceFulfillmentPercent'])}. It filled 1,488 extra units, all in fragmented-supply scenarios.")
p('The safety caps belong to the evaluator, not the application. Fulfillment after caps must not be described as the raw engine executing safely. The comparator is a feasible heuristic, not an optimal or expert-approved policy. No-supply plans can pass feasibility while fulfilling no demand.','Caution')

page('Validity, limitations and research wording')
h('Relationship to the actual studies',2)
p('The synthetic data is related to the studies through average usage scale only. It is not closely calibrated to actual daily behavior. Reported standard deviations, autocorrelation, real stock movements, expiry distributions, clinical mix and supply-delay distributions were not fitted. A fixed seed provides reproducibility, not realism. Generated fields must never be described as recovered hospital observations.')
h('Scope of conclusions',2)
p('The published India data mixes packed red cells and whole blood; Korean usage includes wasted RBC units and lacks RhD labels. Supported whole-blood test surrogates and unspecified/unsigned blood groups do not validate clinical component or blood-group compatibility. The benchmark does not demonstrate transfer safety in actual patient care.')
p('Stock snapshots are independent at each origin. Transfers are assumed immediate and cost-free. Transport constraints, crossmatching, future recipient wastage, replenishment optimization and donor recommendations are not evaluated. The 50% feasibility score depends directly on the deliberately balanced stress-case mix.')
p('Overlapping history windows and shared nominal rate anchors mean forecast cases are not independent. Three seeds assess limited random variation, not uncertainty over all plausible hospitals or modeling assumptions. No population confidence interval or clinical certification is claimed.')
h('Variation across seeds',2)
table(['Seed','Demand WAPE','Shortage accuracy','Plan feasibility'],[
 [key,pct(v['demand']['WAPE_percent']),pct(v['shortage']['accuracyPercent']),pct(v['prescriptive']['feasiblePlanPercent'])] for key,v in S['bySeed'].items()
],[1800,2520,2520,2520])
p('Each seed contributes 935 forecasts and 180 transfer scenarios. These are reproducibility/sensitivity summaries, not confidence limits.','TableSource')
h('Suggested capstone wording',2)
p('BloodConnect was evaluated using controlled synthetic scenarios informed by published aggregate blood usage averages. Across 2,805 seven-day forecasts, MAE was 9.51 units and WAPE was 23.76%; 55.33% fell within 20% of simulated usage. Shortage detection achieved 81.89% accuracy and 85.70% F1 under no replenishment. Of 540 constructed transfer plans, 50% satisfied all tested constraints. Results are exploratory and conditional on uncalibrated simulation assumptions; they do not establish real-world predictive accuracy or clinical effectiveness.')

page('Audit trail, reproduction and references')
h('Reproduction',2)
p('From the BloodConnect repository root, run the following commands separately. They exercise the local engine; the hybrid evaluator does not connect to a database.','Normal')
p('node backend/scripts/testPaperBasedAnalytics.mjs\nnode backend/scripts/evaluateHybridAnalytics.mjs','Small')
p('The paper-based command intentionally exits 1 while application checks fail. A successful hybrid-evaluation exit means its evaluator and internal checks completed, not that the application met an acceptance threshold. Saved synthetic data and scores were reproduced identically on rerun.')
h('Evidence files',2)
for line in [
 'Initial fixture results: backend/simulation/output/validation-report.json',
 'Published fixture values: backend/simulation/paperPublishedData.mjs',
 'Paper-based outputs: backend/simulation/output/paper-based/results.json',
 'Generator: backend/simulation/generateHybridEvaluation.mjs',
 'Independent scoring/replay: backend/simulation/hybridEvaluationMetrics.mjs',
 'Hybrid outputs: backend/simulation/output/hybrid-evaluation/',
 'Within that folder: synthetic-dataset.json, scores.json, case-results.json',
]: p(line,'Small')
p('Internal checks covered nonnegative integer records, history/future separation, seven-day horizons, deterministic generation, hand-calculable forecast/classification metrics, inventory expiry and receipt behavior. Additional saved-artifact checks recomputed MAE, feasible-plan counts, shortage correctness counts and the dataset hash. Syntax checks passed. No engine fix was made for these evaluation runs.')
p('Dataset SHA-256 (compact JSON before final newline):\n'+S['datasetHash'],'Small')
p('Tested analyticsEngine.js SHA-256:\n'+S['sourceHashes']['../../frontend/src/admin/analyticsEngine.js'],'Small')
h('References',2)
p('[1] Bansal et al. (2023). Blood Inventory Management During COVID-19 Pandemic Using a Simple Mathematical Tool: A Two-Year Study from a Tertiary Care Hospital in North India. Indian Journal of Hematology and Blood Transfusion, 39, 655-661. https://doi.org/10.1007/s12288-023-01631-8','Small')
p('[2] Park, Lim and Kim (2024; online 2023). Optimizing the Hospital Blood Bank Stock in Korea: A Comparative Analysis of the Uniform 5-Day Stock Index and a Novel Blood Stock Index. Annals of Laboratory Medicine, 44, 262-270. https://doi.org/10.3343/alm.2023.0242','Small')
p('[3] Hyndman and Athanasopoulos. Forecasting: Principles and Practice, 3rd ed., section 5.8, Evaluating point forecast accuracy. https://otexts.com/fpp3/accuracy.html','Small')
p('[4] scikit-learn documentation. Metrics and scoring: quantifying the quality of predictions. https://scikit-learn.org/stable/modules/model_evaluation.html','Small')

# Structural QA: preset geometry and every table's explicit widths.
assert section.page_width.twips == 12240 and section.page_height.twips == 15840
assert all(x.twips == 1440 for x in [section.top_margin,section.bottom_margin,section.left_margin,section.right_margin])
for t in doc.tables:
    widths = [int(c.get(qn('w:w'))) for c in t._tbl.tblGrid]
    assert sum(widths) == 9360
    for row in t.rows:
        assert [int(c._tc.get_or_add_tcPr().find(qn('w:tcW')).get(qn('w:w'))) for c in row.cells] == widths
    assert t.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None
text = '\n'.join(x.text for x in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
assert '55.33%' in text and '81.89%' in text and '50.00%' in text and 'not closely calibrated' in text
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(json.dumps({'path':str(OUT),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs),'structural_qa':'passed'}))
